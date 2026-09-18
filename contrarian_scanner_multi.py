import html
import json
import logging
import os
import smtplib
import sys
import time
from datetime import datetime, timezone
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from io import BytesIO

import pandas as pd
import pandas_market_calendars as mcal
import requests
import yfinance as yf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ============================================================
# Configuratie
# ============================================================
DEEPSEEK_API_KEY = os.getenv("DEEPSEEK_API_KEY")
SENDER_EMAIL = os.getenv("SENDER_EMAIL")
SENDER_PASSWORD = os.getenv("SENDER_PASSWORD")
RECEIVER_EMAIL = os.getenv("RECEIVER_EMAIL")
TOP_N = int(os.getenv("TOP_N", "5"))
PERIOD = os.getenv("PERIOD", "1y")
MIN_HISTORY = 200
REQUEST_TIMEOUT = int(os.getenv("REQUEST_TIMEOUT", "90"))

# Drempel waarboven/waaronder een netto contraire score als Koop/Verkoop telt
SIGNAAL_DREMPEL = 2.0

# Buffer (in aantal ATR's) die boven op de recente 10-daagse low/high wordt
# gelegd om de stop-loss te bepalen. Groter = ruimere stop, minder kans om
# uitgeschud te worden door normale ruis, maar wel meer risico per trade.
ATR_STOP_BUFFER = 1.0

# Retry-instellingen voor de yfinance-download
DOWNLOAD_MAX_POGINGEN = 3
DOWNLOAD_BACKOFF_SECONDEN = 5  # verdubbelt elke poging

# ---- Drie aparte universes ----
AEX_TICKERS = [
    "ASML.AS", "SHELL.AS", "UNA.AS", "PRX.AS", "INGA.AS", "REN.AS",
    "ASM.AS", "MT.AS", "HEIA.AS", "AD.AS", "ADYEN.AS", "ABN.AS",
    "PHIA.AS", "NN.AS", "BESI.AS", "KPN.AS", "WKL.AS", "ASRNL.AS",
    "AGN.AS", "DSFIR.AS", "UMG.AS", "EXO.AS", "SBMO.AS", "IMCD.AS",
    "AALB.AS",
]

DOW_TICKERS = [
    "AAPL", "AMGN", "AMZN", "AXP", "BA", "CAT", "CRM", "CSCO", "CVX", "DIS",
    "GS", "HD", "HON", "IBM", "JNJ", "JPM", "KO", "MCD", "MMM", "MRK",
    "MSFT", "NKE", "NVDA", "PG", "SHW", "TRV", "UNH", "V", "VZ", "WMT"
]

NASDAQ_TICKERS = [
    "AAPL", "MSFT", "NVDA", "AMZN", "GOOGL", "GOOG", "META", "TSLA", "AVGO", "AMD",
    "INTC", "QCOM", "TXN", "AMAT", "LRCX", "KLAC", "ADI", "MU", "ADBE", "ORCL",
    "NOW", "PANW", "NFLX", "COST", "PEP", "SBUX", "BKNG", "ABNB", "PYPL", "COIN",
    "SHOP", "UBER", "CRM", "CSCO"
]

# (naam, tickerlijst, valutasymbool)
UNIVERSES = [
    ("AEX", AEX_TICKERS, "€"),
    ("Dow", DOW_TICKERS, "$"),
    ("Nasdaq", NASDAQ_TICKERS, "$"),
]

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)


class DataOphalenMislukt(Exception):
    """Wordt gegooid als de koersdata na alle retries niet opgehaald kon worden."""
    pass


def is_trading_day(calendar_code, date=None):
    date = date or datetime.now(timezone.utc).date()
    return not mcal.get_calendar(calendar_code).schedule(start_date=date, end_date=date).empty


def get_ticker_frame(data, ticker, only_one_ticker):
    try:
        if only_one_ticker or not isinstance(data.columns, pd.MultiIndex):
            df = data.copy()
        else:
            if ticker not in data.columns.get_level_values(0):
                return None
            df = data[ticker].copy()
        df = df.dropna(how="all")
        if df.empty:
            return None
        df.columns = [str(c).lower() for c in df.columns]
        if "adj close" in df.columns and "close" not in df.columns:
            df["close"] = df["adj close"]
        return df
    except Exception:
        return None


def add_indicators(df):
    """
    Indicatoren gericht op CONTRAIR beleggen: statistische extremen en
    sentiment-uitschieters, plus ATR en recente 10-daagse hoog/laag om
    realistische koersdoelen en stop-loss-niveaus te kunnen berekenen.
    """
    df = df.copy()

    # ---- RSI (Wilder's smoothing) ----
    delta = df["close"].diff()
    gain = delta.where(delta > 0, 0.0)
    loss = -delta.where(delta < 0, 0.0)
    avg_gain = gain.ewm(alpha=1 / 14, min_periods=14, adjust=False).mean()
    avg_loss = loss.ewm(alpha=1 / 14, min_periods=14, adjust=False).mean()
    rs = avg_gain / avg_loss
    df["rsi_14"] = 100 - (100 / (1 + rs))

    # ---- Bollinger Bands (20, 2 stddev) ----
    df["sma_20"] = df["close"].rolling(20).mean()
    df["std_20"] = df["close"].rolling(20).std()
    df["bb_upper"] = df["sma_20"] + 2 * df["std_20"]
    df["bb_lower"] = df["sma_20"] - 2 * df["std_20"]
    df["percent_b"] = (df["close"] - df["bb_lower"]) / (df["bb_upper"] - df["bb_lower"])

    # ---- Z-score t.o.v. 50-daags gemiddelde ----
    df["sma_50"] = df["close"].rolling(50).mean()
    df["std_50"] = df["close"].rolling(50).std()
    df["z_score_50"] = (df["close"] - df["sma_50"]) / df["std_50"]

    # ---- 52-weken hoogte/laagte ----
    df["hoog_52w"] = df["high"].rolling(252, min_periods=100).max()
    df["laag_52w"] = df["low"].rolling(252, min_periods=100).min()
    df["afstand_tot_laag_52w_pct"] = (df["close"] - df["laag_52w"]) / df["laag_52w"] * 100
    df["afstand_tot_hoog_52w_pct"] = (df["hoog_52w"] - df["close"]) / df["hoog_52w"] * 100

    # ---- Volume ----
    df["vol_sma"] = df["volume"].rolling(20).mean()
    df["rel_volume"] = df["volume"] / df["vol_sma"]
    df["dagrendement_pct"] = df["close"].pct_change() * 100

    # ---- ATR (Average True Range, Wilder's smoothing) ----
    # Gebruikt om een realistische stop-loss-afstand te bepalen die rekening
    # houdt met de actuele volatiliteit van het aandeel, i.p.v. een vast
    # percentage dat voor een rustig en een wild aandeel even groot zou zijn.
    prev_close = df["close"].shift(1)
    true_range = pd.concat([
        df["high"] - df["low"],
        (df["high"] - prev_close).abs(),
        (df["low"] - prev_close).abs(),
    ], axis=1).max(axis=1)
    df["atr_14"] = true_range.ewm(alpha=1 / 14, min_periods=14, adjust=False).mean()

    # ---- Recente 10-daagse hoog/laag ----
    # Basis voor de stop-loss: net voorbij het recente swing-punt, plus een
    # ATR-buffer, zodat normale ruis de positie niet meteen uitschudt.
    df["laag_10d"] = df["low"].rolling(10).min()
    df["hoog_10d"] = df["high"].rolling(10).max()

    return df


def _bereken_target_en_stop(latest, direction):
    """
    Berekent een realistisch koersdoel en stop-loss voor een contraire
    trade, op basis van mean-reversion (doel = 20-daags gemiddelde) en
    ATR-gecorrigeerde swing-niveaus (stop-loss).
    """
    entry = latest.get("close")
    sma_20 = latest.get("sma_20")
    atr = latest.get("atr_14")

    if direction == "Koop":
        laag_10d = latest.get("laag_10d")
        if pd.isna(entry) or pd.isna(sma_20) or pd.isna(atr) or pd.isna(laag_10d):
            return "N/A", "N/A", "N/A"
        target = sma_20  # reversie naar het korte-termijn gemiddelde
        stop = min(laag_10d, entry) - ATR_STOP_BUFFER * atr
        if stop >= entry or target <= entry:
            return "N/A", "N/A", "N/A"
        risico = entry - stop
        beloning = target - entry
        rr = round(beloning / risico, 2) if risico > 0 else "N/A"
        return round(float(target), 2), round(float(stop), 2), rr

    elif direction == "Verkoop":
        hoog_10d = latest.get("hoog_10d")
        if pd.isna(entry) or pd.isna(sma_20) or pd.isna(atr) or pd.isna(hoog_10d):
            return "N/A", "N/A", "N/A"
        target = sma_20  # reversie naar het korte-termijn gemiddelde
        stop = max(hoog_10d, entry) + ATR_STOP_BUFFER * atr
        if stop <= entry or target >= entry:
            return "N/A", "N/A", "N/A"
        risico = stop - entry
        beloning = entry - target
        rr = round(beloning / risico, 2) if risico > 0 else "N/A"
        return round(float(target), 2), round(float(stop), 2), rr

    return "N/A", "N/A", "N/A"


def score_signal(df):
    """
    Contraire scoring: beloont extremen tegen de heersende stemming in.
    Retourneert ook een berekend koersdoel, stop-loss en risk/reward —
    geen AI-schatting, maar direct afgeleid uit ATR en 20-daags gemiddelde.
    """
    if len(df) < 60:
        return None

    latest = df.iloc[-1]
    score = 0.0
    reasons = []

    rsi = latest.get("rsi_14")
    if pd.isna(rsi):
        rsi = 50

    if rsi < 25:
        score += 2.2
        reasons.append(f"RSI extreem oversold ({rsi:.1f}) — mogelijke capitulatie")
    elif rsi < 32:
        score += 1.0
        reasons.append(f"RSI oversold ({rsi:.1f})")
    elif rsi > 75:
        score -= 2.2
        reasons.append(f"RSI extreem overbought ({rsi:.1f}) — mogelijke euforie-top")
    elif rsi > 68:
        score -= 1.0
        reasons.append(f"RSI overbought ({rsi:.1f})")

    percent_b = latest.get("percent_b")
    if pd.notna(percent_b):
        if percent_b < 0:
            score += 1.5
            reasons.append("Koers doorbreekt de onderkant van de Bollinger Band")
        elif percent_b > 1:
            score -= 1.5
            reasons.append("Koers doorbreekt de bovenkant van de Bollinger Band")

    z_score = latest.get("z_score_50")
    if pd.notna(z_score):
        if z_score <= -2.0:
            score += 1.8
            reasons.append(f"Koers {abs(z_score):.1f}σ onder 50-daags gemiddelde (statistische uitschieter)")
        elif z_score >= 2.0:
            score -= 1.8
            reasons.append(f"Koers {z_score:.1f}σ boven 50-daags gemiddelde (statistische uitschieter)")

    afstand_laag = latest.get("afstand_tot_laag_52w_pct")
    afstand_hoog = latest.get("afstand_tot_hoog_52w_pct")
    if pd.notna(afstand_laag) and afstand_laag <= 3:
        score += 1.0
        reasons.append("Koers nabij 52-weken dieptepunt")
    if pd.notna(afstand_hoog) and afstand_hoog <= 3:
        score -= 1.0
        reasons.append("Koers nabij 52-weken hoogtepunt")

    rel_vol = latest.get("rel_volume")
    dagrendement = latest.get("dagrendement_pct")
    if pd.notna(rel_vol) and pd.notna(dagrendement) and rel_vol > 1.8:
        if dagrendement < -3:
            score += 0.8
            reasons.append("Paniekvolume op een sterk dalende dag (mogelijke uitverkoop-bodem)")
        elif dagrendement > 3:
            score -= 0.8
            reasons.append("Uitzonderlijk volume op een sterk stijgende dag (mogelijke euforie-top)")

    if score >= SIGNAAL_DREMPEL:
        direction = "Koop"
    elif score <= -SIGNAAL_DREMPEL:
        direction = "Verkoop"
    else:
        direction = "Neutraal"

    target, stop, rr = _bereken_target_en_stop(latest, direction)

    return {
        "Contrarian Score": round(score, 2),
        "Direction": direction,
        "RSI": round(float(rsi), 1),
        "%B (Bollinger)": round(float(percent_b), 2) if pd.notna(percent_b) else "N/A",
        "Z-score(50)": round(float(z_score), 2) if pd.notna(z_score) else "N/A",
        "% tot 52w-laag": round(float(afstand_laag), 1) if pd.notna(afstand_laag) else "N/A",
        "% tot 52w-hoog": round(float(afstand_hoog), 1) if pd.notna(afstand_hoog) else "N/A",
        "Rel. Volume": round(float(rel_vol), 2) if pd.notna(rel_vol) else "N/A",
        "Koersdoel": target,
        "Stop-loss": stop,
        "Risk/Reward": rr,
        "Reasons": " | ".join(reasons) if reasons else "Geen extreme afwijking",
    }


def safe_round(val, digits=2):
    try:
        return round(float(val), digits)
    except Exception:
        return "N/A"


def download_met_retry(tickers):
    laatste_fout = None
    wachttijd = DOWNLOAD_BACKOFF_SECONDEN

    for poging in range(1, DOWNLOAD_MAX_POGINGEN + 1):
        try:
            data = yf.download(
                tickers, period=PERIOD, interval="1d", group_by="ticker",
                auto_adjust=False, threads=True, progress=False,
            )
            if data is None or data.empty:
                raise ValueError("yfinance gaf een lege dataset terug")
            return data
        except Exception as e:
            laatste_fout = e
            logger.warning(f"Download poging {poging}/{DOWNLOAD_MAX_POGINGEN} mislukt: {e}")
            if poging < DOWNLOAD_MAX_POGINGEN:
                time.sleep(wachttijd)
                wachttijd *= 2

    raise DataOphalenMislukt(f"Download definitief mislukt na {DOWNLOAD_MAX_POGINGEN} pogingen: {laatste_fout}")


def process_universe(tickers, data, universe_name, currency_symbol, only_one_ticker):
    results = []
    for ticker in tickers:
        try:
            df = get_ticker_frame(data, ticker, only_one_ticker)
            if df is None or len(df) < MIN_HISTORY:
                continue
            df = add_indicators(df.reset_index(drop=True))
            signal = score_signal(df)
            if not signal:
                continue
            latest = df.iloc[-1]
            results.append({
                "Ticker": ticker,
                "Universe": universe_name,
                "Valuta": currency_symbol,
                "Koers": safe_round(latest.get("close")),
                **signal
            })
        except Exception as e:
            logger.warning(f"Fout bij {ticker}: {e}")

    buys = [r for r in results if r["Direction"] == "Koop"]
    sells = [r for r in results if r["Direction"] == "Verkoop"]
    buys.sort(key=lambda x: x["Contrarian Score"], reverse=True)
    sells.sort(key=lambda x: x["Contrarian Score"])
    return buys[:TOP_N], sells[:TOP_N]


def fetch_all():
    alle_tickers = list(set(t for _, tickers, _ in UNIVERSES for t in tickers))
    logger.info(f"Data ophalen voor {len(alle_tickers)} unieke tickers over {len(UNIVERSES)} universes...")
    data = download_met_retry(alle_tickers)  # gooit DataOphalenMislukt bij falen
    only_one_ticker = len(alle_tickers) == 1

    resultaten = {}
    for naam, tickers, valuta in UNIVERSES:
        buys, sells = process_universe(tickers, data, naam, valuta, only_one_ticker)
        resultaten[f"{naam.lower()}_buys"] = buys
        resultaten[f"{naam.lower()}_sells"] = sells
        logger.info(f"{naam:8s} → {len(buys)} contraire koop / {len(sells)} contraire verkoop")

    return resultaten


def analyze_with_ai(resultaten):
    if not DEEPSEEK_API_KEY:
        return "Geen AI-analyse (API key ontbreekt)."

    system_prompt = (
        "Je bent een ervaren contraire belegger ('contrarian investor') actief op de "
        "AEX, Dow 30 en Nasdaq. Je zoekt bewust extremen tegen de heersende "
        "marktstemming in: paniek en capitulatie als koopkans, euforie en overdreven "
        "optimisme als verkoopkans. Elk signaal bevat al een berekend koersdoel, "
        "stop-loss en risk/reward — reken deze niet zelf opnieuw uit, maar interpreteer "
        "en beoordeel ze. Je antwoordt uitsluitend in het Nederlands, volgt exact de "
        "gevraagde structuur, en waarschuwt expliciet dat een contraire positie tegen "
        "de trend in kan lopen ('probeer geen vallend mes te vangen')."
    )

    prompt = f"""
Hier zijn de contraire signalen voor drie universes. Elk signaal bevat al een
berekend koersdoel (reversie naar 20-daags gemiddelde), stop-loss (ATR-gecorrigeerd
t.o.v. het recente 10-daagse swing-niveau) en Risk/Reward-ratio.

=== AEX ===
Koopkansen: {json.dumps(resultaten['aex_buys'], indent=2, ensure_ascii=False)}
Verkoopkansen: {json.dumps(resultaten['aex_sells'], indent=2, ensure_ascii=False)}

=== DOW 30 ===
Koopkansen: {json.dumps(resultaten['dow_buys'], indent=2, ensure_ascii=False)}
Verkoopkansen: {json.dumps(resultaten['dow_sells'], indent=2, ensure_ascii=False)}

=== NASDAQ ===
Koopkansen: {json.dumps(resultaten['nasdaq_buys'], indent=2, ensure_ascii=False)}
Verkoopkansen: {json.dumps(resultaten['nasdaq_sells'], indent=2, ensure_ascii=False)}

Schrijf een gestructureerd rapport in het Nederlands met EXACT deze opbouw:

## 1. AEX – Beste Contraire Kansen
De 1-2 sterkste koop- én verkoopsignalen: ticker, koers, status (KOPEN/VERKOPEN,
contrair), waarom dit overdreven pessimisme/optimisme lijkt, en een beoordeling
van het gegeven koersdoel/stop-loss/risk-reward (is dit realistisch, gezien de
onderliggende reden voor de beweging?).

## 2. Dow 30 – Beste Contraire Kansen
Zelfde structuur.

## 3. Nasdaq – Beste Contraire Kansen
Zelfde structuur.

## 4. Marktobservatie
Eén korte observatie: waar zitten de meeste extremen vandaag (AEX, Dow of Nasdaq), en wat zegt dat over de marktbrede stemming?

## 5. Waarschuwing
Een korte, nuchtere waarschuwing over de risico's van contrair beleggen.

BELANGRIJK:
- Geef NOOIT tegenstrijdige adviezen (een aandeel mag niet tegelijk koop én verkoop zijn).
- Herbereken koersdoel/stop-loss niet zelf — beoordeel de gegeven waarden.
- Wees nuchter en realistisch. Geen garanties.
"""

    try:
        resp = requests.post(
            "https://api.deepseek.com/chat/completions",
            json={
                "model": "deepseek-chat",
                "messages": [
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": prompt},
                ],
                "temperature": 0.15,
            },
            headers={
                "Authorization": f"Bearer {DEEPSEEK_API_KEY}",
                "Content-Type": "application/json",
            },
            timeout=REQUEST_TIMEOUT,
        )
        resp.raise_for_status()
        return resp.json()["choices"][0]["message"]["content"]
    except Exception as e:
        logger.error(f"AI mislukt: {e}")
        return f"AI-analyse mislukt: {e}"


def _stijl_tabel_header(table, achtergrondkleur):
    for cell in table.rows[0].cells:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), achtergrondkleur)
        cell._tc.get_or_add_tcPr().append(shading)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(9)


def _voeg_signalen_tabel_toe(doc, titel, resultaten, geen_data_tekst, header_kleur):
    doc.add_heading(titel, level=2)
    if not resultaten:
        p = doc.add_paragraph(geen_data_tekst)
        p.italic = True
        return

    headers = ["Ticker", "Koers", "RSI", "%B", "Z-score", "Koersdoel", "Stop-loss", "R/R", "Score"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    _stijl_tabel_header(table, header_kleur)

    for r in resultaten:
        row = table.add_row().cells
        row[0].text = r["Ticker"]
        row[1].text = f"{r['Valuta']}{r['Koers']}"
        row[2].text = str(r["RSI"])
        row[3].text = str(r["%B (Bollinger)"])
        row[4].text = str(r["Z-score(50)"])
        row[5].text = f"{r['Valuta']}{r['Koersdoel']}" if r["Koersdoel"] != "N/A" else "N/A"
        row[6].text = f"{r['Valuta']}{r['Stop-loss']}" if r["Stop-loss"] != "N/A" else "N/A"
        row[7].text = str(r["Risk/Reward"])
        row[8].text = str(r["Contrarian Score"])


def create_docx_report(resultaten, ai_report):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1)
    section.bottom_margin = Cm(1)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    title = doc.add_heading(f"Contrarian Scanner (AEX / Dow / Nasdaq) – {datetime.now():%d-%m-%Y}", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Gegenereerd: {datetime.now():%d-%m-%Y %H:%M UTC}")
    doc.add_paragraph(
        "Focus: contraire signalen met berekend koersdoel (reversie naar 20-daags "
        "gemiddelde) en ATR-gecorrigeerde stop-loss"
    )

    for naam in ["AEX", "Dow", "Nasdaq"]:
        key = naam.lower()
        _voeg_signalen_tabel_toe(doc, f"{naam} – Contraire Koopkansen (overdreven pessimisme)",
                                  resultaten[f"{key}_buys"],
                                  f"Geen extreme oversold-signalen gevonden voor {naam} vandaag.", "2E7D32")
        _voeg_signalen_tabel_toe(doc, f"{naam} – Contraire Verkoopkansen (overdreven optimisme)",
                                  resultaten[f"{key}_sells"],
                                  f"Geen extreme overbought-signalen gevonden voor {naam} vandaag.", "C62828")

    doc.add_heading("DeepSeek Analyse – Contraire Kansen + Risicowaarschuwing", level=2)
    doc.add_paragraph(ai_report)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def send_email(resultaten, ai_report):
    if not all([SENDER_EMAIL, SENDER_PASSWORD, RECEIVER_EMAIL]):
        logger.error("E-mailgegevens ontbreken (SENDER_EMAIL / SENDER_PASSWORD / RECEIVER_EMAIL)")
        return False

    receivers = [e.strip() for e in RECEIVER_EMAIL.split(",") if e.strip()]
    if not receivers:
        logger.error("RECEIVER_EMAIL bevat geen geldige adressen")
        return False

    msg = MIMEMultipart("mixed")
    msg["Subject"] = f"Contrarian Scanner (AEX/Dow/Nasdaq) – {datetime.now():%d-%m-%Y}"
    msg["From"] = SENDER_EMAIL
    msg["To"] = ", ".join(receivers)

    def maak_tabel(data):
        if not data:
            return "<p>Geen extreme signalen</p>"
        df = pd.DataFrame(data)
        df["Koers"] = df.apply(lambda r: f"{r['Valuta']}{r['Koers']}", axis=1)
        df["Koersdoel"] = df.apply(lambda r: f"{r['Valuta']}{r['Koersdoel']}" if r["Koersdoel"] != "N/A" else "N/A", axis=1)
        df["Stop-loss"] = df.apply(lambda r: f"{r['Valuta']}{r['Stop-loss']}" if r["Stop-loss"] != "N/A" else "N/A", axis=1)
        kolommen = ["Ticker", "Koers", "RSI", "Koersdoel", "Stop-loss", "Risk/Reward", "Contrarian Score", "Reasons"]
        return df[kolommen].to_html(index=False, escape=True)

    secties = ""
    for naam in ["AEX", "Dow", "Nasdaq"]:
        key = naam.lower()
        secties += f"""
        <h3 style="color:green">{naam} – Contraire Koopkansen</h3>
        {maak_tabel(resultaten[f'{key}_buys'])}
        <h3 style="color:red">{naam} – Contraire Verkoopkansen</h3>
        {maak_tabel(resultaten[f'{key}_sells'])}
        """

    body = f"""
    <html><body style="font-family:Arial,sans-serif;line-height:1.5">
    <h2>Contrarian Scanner (AEX / Dow / Nasdaq) – {datetime.now():%d-%m-%Y}</h2>
    {secties}
    <h3>AI Analyse + Risicowaarschuwing</h3>
    <div>{html.escape(ai_report).replace(chr(10), "<br>")}</div>

    <p style="color:#666;font-size:12px;margin-top:25px">
    Dit is geen beleggingsadvies. Koersdoel en stop-loss zijn statistisch afgeleid
    (20-daags gemiddelde / ATR) en geen garantie. Contrair beleggen kent verhoogd
    risico: een extreem kan langer aanhouden of verder doorzetten dan verwacht.
    </p>
    </body></html>
    """
    msg.attach(MIMEText(body, "html"))

    try:
        attachment = MIMEBase("application", "vnd.openxmlformats-officedocument.wordprocessingml.document")
        attachment.set_payload(create_docx_report(resultaten, ai_report).read())
        encoders.encode_base64(attachment)
        attachment.add_header("Content-Disposition", "attachment",
                              filename=f"Contrarian_Scanner_{datetime.now():%Y%m%d}.docx")
        msg.attach(attachment)
    except Exception as e:
        logger.warning(f"DOCX mislukt: {e}")

    try:
        with smtplib.SMTP("smtp.gmail.com", 587, timeout=30) as server:
            server.starttls()
            server.login(SENDER_EMAIL, SENDER_PASSWORD)
            server.sendmail(SENDER_EMAIL, receivers, msg.as_string())
        logger.info(f"E-mail verzonden naar: {', '.join(receivers)}")
        return True
    except Exception as e:
        logger.error(f"E-mail mislukt: {e}")
        return False


def main():
    logger.info("=" * 60)
    logger.info("START Contrarian Scanner (AEX / Dow / Nasdaq)")
    logger.info("=" * 60)

    is_manual = os.getenv("GITHUB_EVENT_NAME") == "workflow_dispatch"
    if not is_manual:
        aex_open = is_trading_day("XAMS")
        us_open = is_trading_day("NYSE")
        if not aex_open and not us_open:
            logger.info("Geen enkele relevante beurs open vandaag – scan overgeslagen")
            sys.exit(0)

    try:
        resultaten = fetch_all()
    except DataOphalenMislukt as e:
        logger.critical(f"Kon geen koersdata ophalen: {e}")
        sys.exit(1)

    if not any(resultaten.values()):
        # Bij contrair beleggen is "geen extremen" een normale, verwachte
        # uitkomst op een rustige dag over drie markten heen — geen fout.
        logger.info("Geen extreme afwijkingen gevonden in AEX, Dow of Nasdaq — rustige marktdag")
        sys.exit(0)

    ai_report = analyze_with_ai(resultaten)

    output = os.getenv("OUTPUT_DIR", "output")
    os.makedirs(output, exist_ok=True)
    with open(f"{output}/contrarian_results.json", "w", encoding="utf-8") as f:
        json.dump(resultaten, f, indent=2, ensure_ascii=False)
    with open(f"{output}/contrarian_ai_report.txt", "w", encoding="utf-8") as f:
        f.write(ai_report)

    email_verzonden = send_email(resultaten, ai_report)
    if not email_verzonden:
        logger.critical("Analyse geslaagd, maar e-mail kon niet worden verzonden")
        sys.exit(1)

    logger.info("SCRIPT VOLLEDIG AFGEROND")


if __name__ == "__main__":
    main()
